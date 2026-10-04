"""Builds the Equity Value Creation Bridge Toolkit workbook (Post Fiat Analytical Infrastructure Library).
Usage: python build_value_creation_bridge.py [out.xlsx]"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

OUT = sys.argv[1] if len(sys.argv) > 1 else "Equity_Value_Creation_Bridge_Toolkit.xlsx"
N = ["ReadMe - Validator Instructions", "Entry and Exit Assumptions", "Value Driver Inputs",
     "Equity Value Creation Bridge", "Lever Impact Ranking", "Summary Recommendation",
     "Anonymized Worked Example"]
wb = Workbook(); wb.active.title = N[0]
for n in N[1:]:
    wb.create_sheet(n)
RM, EA, VD, BR, LR, SR, WE = (wb[n] for n in N)
A, V, B, L, S, W = (f"'{n}'!" for n in N[1:])

thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"in": PatternFill("solid", fgColor="FFF2CC"), "f": PatternFill("solid", fgColor="DDEBF7"),
        "h": PatternFill("solid", fgColor="2F5597"), "t": PatternFill("solid", fgColor="1F3864"),
        "k": PatternFill("solid", fgColor="E2EFDA")}
FMT = {"$M": '#,##0.00', "%": '0.0%', "x": '0.00"x"', "n": '0', "yrs": '0', "bps": '#,##0'}


def title(ws, text, sub, span):
    for c in range(1, span + 1):
        ws.cell(1, c).fill = FILL["t"]
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")


def header(ws, r, cols):
    for i, v in enumerate(cols, 1):
        c = ws.cell(r, i, v); c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL["h"]
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX


def put(ws, r, col, v, fmt=None, kind=None, bold=False):
    c = ws.cell(r, col, v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind is None and isinstance(v, str) and v.startswith("="):
        kind = "f"
    if kind:
        c.fill = FILL[kind]
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if bold:
        c.font = Font(bold=True)
    return c


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def section(ws, r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=12, color="1F3864")


def table(ws, start, items):
    """items: (label, value, unit, note). Numeric constants = yellow inputs; formulas = blue."""
    for i, (lab, val, unit, note) in enumerate(items):
        r = start + i
        isf = isinstance(val, str) and val.startswith("=")
        kind = "f" if isf else (None if isinstance(val, str) else "in")
        put(ws, r, 1, lab, bold=True); put(ws, r, 2, val, unit if unit in FMT else None, kind)
        put(ws, r, 3, unit); put(ws, r, 4, note)


# ---------------- Entry and Exit Assumptions ----------------
title(EA, "Entry and Exit Assumptions", "Yellow = input, blue = formula. $M unless noted. Change inputs here and on Value Driver Inputs; every other tab recalculates.", 4)
header(EA, 3, ["Input", "Value", "Unit", "Source / note"])
table(EA, 4, [
    ("Deal code name", "HoldCo B (anonymized)", "text", "Code name only; no company identifiers."),
    ("Entry LTM revenue", 60.0, "$M", "Scaled, illustrative LMM services platform."),
    ("Entry EBITDA margin", 0.15, "%", "Quality-of-earnings adjusted."),
    ("Entry LTM EBITDA", "=B5*B6", "$M", "Revenue x margin."),
    ("Entry EV / EBITDA multiple", 8.0, "x", "Purchase multiple paid."),
    ("Entry enterprise value", "=B7*B8", "$M", "EBITDA x multiple."),
    ("Entry net debt", 40.0, "$M", "Acquisition debt less closing cash."),
    ("Entry leverage (net debt / EBITDA)", "=B10/B7", "x", "Sanity check vs lender appetite."),
    ("Entry equity value", "=B9-B10", "$M", "Sponsor equity cheque (fees ignored)."),
    ("Hold period", 5, "yrs", "Whole years, 1-7 (schedule supports up to 7)."),
    ("Underwritten exit EV / EBITDA multiple", 8.0, "x", "Judgment call JC1: held flat at entry multiple."),
    ("Management-case exit multiple (reference only)", 9.5, "x", "Not underwritten; sizes JC1."),
    ("Target gross MOIC", 2.5, "x", "Fund underwriting hurdle."),
    ("Target gross IRR", 0.20, "%", "Fund underwriting hurdle."),
    ("Base-case MOIC floor", 1.75, "x", "Below this, do not underwrite."),
    ("Downside-case MOIC floor", 1.25, "x", "Capital-preservation test."),
    ("Max share of value created from multiple change", 0.25, "%", "Limits reliance on market re-rating."),
    ("Min share of value created from operational levers", 0.50, "%", "Organic growth + add-on + margin."),
    ("Downside: exit multiple haircut", 1.0, "x", "Turns below underwritten exit multiple."),
    ("Downside: organic growth credited", 0.5, "%", "Share of base organic CAGR achieved."),
    ("Downside: margin change credited", 0.0, "%", "Share of base margin change achieved."),
])
widths(EA, [50, 24, 7, 58]); EA.freeze_panes = "A4"

# ---------------- Value Driver Inputs ----------------
title(VD, "Value Driver Inputs", "Operating and financial levers over the hold period, plus the annual EBITDA / free cash flow schedule.", 6)
header(VD, 3, ["Driver", "Value", "Unit", "Source / note"])
table(VD, 4, [
    ("Organic revenue CAGR", 0.06, "%", "Price + volume; excludes add-ons."),
    ("Organic exit revenue", f"={A}B5*(1+B4)^{A}B13", "$M", "Entry revenue compounded over the hold."),
    ("Add-on revenue acquired (run-rate at exit)", 12.0, "$M", "JC3: one signed-LOI add-on only."),
    ("Add-on EBITDA margin at acquisition", 0.12, "%", "Pre-synergy margin of the add-on."),
    ("Add-on purchase multiple", 5.5, "x", "Multiple arbitrage vs platform entry multiple."),
    ("Add-on purchase price (debt-funded)", "=B6*B7*B8", "$M", "Funded with incremental debt."),
    ("Exit revenue", "=B5+B6", "$M", "Organic + add-on."),
    ("Exit EBITDA margin", 0.17, "%", "JC2: +200 bps credited of a +400 bps plan."),
    ("Exit EBITDA", "=B10*B11", "$M", "Exit revenue x exit margin."),
    ("Margin change vs entry", f"=(B11-{A}B6)*10000", "bps", "Basis points."),
    ("Cash conversion (FCF / EBITDA)", 0.50, "%", "JC4: after capex, cash tax and interest."),
    ("Share of FCF swept to debt paydown", 0.80, "%", "Remainder retained as balance-sheet cash."),
    ("Management-case exit margin (reference only)", 0.19, "%", "Not underwritten; sizes JC2."),
    ("Management-case cash conversion (reference only)", 0.65, "%", "Not underwritten; sizes JC4."),
    ("Add-on pipeline revenue excluded (reference only)", 20.0, "$M", "Two unsigned targets; sizes JC3."),
])
section(VD, 20, "Annual EBITDA and free cash flow schedule (linear EBITDA path from entry to exit)")
header(VD, 21, ["Year", "In hold? (1/0)", "EBITDA ($M)", "Free cash flow ($M)", "Debt paydown ($M)", "Cash retained ($M)"])
for t in range(1, 8):
    r = 21 + t
    put(VD, r, 1, t)
    put(VD, r, 2, f"=IF(A{r}<={A}$B$13,1,0)", "n")
    put(VD, r, 3, f"=B{r}*({A}$B$7+($B$12-{A}$B$7)*A{r}/{A}$B$13)", "$M")
    put(VD, r, 4, f"=C{r}*$B$14", "$M")
    put(VD, r, 5, f"=D{r}*$B$15", "$M")
    put(VD, r, 6, f"=D{r}-E{r}", "$M")
put(VD, 29, 1, "Total", bold=True)
for c in "BCDEF":
    put(VD, 29, ord(c) - 64, f"=SUM({c}22:{c}28)", "n" if c == "B" else "$M", bold=True)
put(VD, 30, 1, "Check: years modelled = hold period", bold=True)
put(VD, 30, 2, f'=IF(B29={A}B13,"OK","ERROR - hold must be 1-7 years")')
widths(VD, [50, 16, 14, 44, 16, 16]); VD.freeze_panes = "A4"

# ---------------- Equity Value Creation Bridge ----------------
title(BR, "Equity Value Creation Bridge", "Entry equity -> exit equity. Every step is a live formula; the additive bridge must reconcile to directly computed exit equity.", 6)
header(BR, 3, ["Step", "Lever", "Logic", "Value ($M)", "% of value created", "Live formula (FORMULATEXT)"])
steps = [
    ("Start", "Entry equity value", "Entry EV less entry net debt", f"={A}B12"),
    ("Step 1", "Organic revenue growth", "(Organic exit revenue - entry revenue) x entry margin x entry multiple", f"=({V}B5-{A}B5)*{A}B6*{A}B8"),
    ("Step 2", "Add-on revenue", "Add-on revenue x entry margin x entry multiple", f"={V}B6*{A}B6*{A}B8"),
    ("Step 3", "Margin change", "Exit revenue x (exit margin - entry margin) x entry multiple", f"={V}B10*({V}B11-{A}B6)*{A}B8"),
    ("Step 4", "Multiple change", "Exit EBITDA x (exit multiple - entry multiple)", f"={V}B12*({A}B14-{A}B8)"),
    ("Step 5", "Add-on purchase price (debt-funded)", "Incremental debt raised to buy the add-on", f"=-{V}B9"),
    ("Step 6", "Debt paydown", "Cumulative FCF swept to debt", f"={V}E29"),
    ("Step 7", "Cash generation (retained)", "Cumulative FCF retained as cash", f"={V}F29"),
    ("End", "Exit equity value (bridge)", "Entry equity + steps 1-7", "=D4+SUM(D5:D11)"),
]
for i, (s, lev, logic, f) in enumerate(steps):
    r = 4 + i; ends = s in ("Start", "End")
    put(BR, r, 1, s, bold=True); put(BR, r, 2, lev, bold=ends); put(BR, r, 3, logic)
    put(BR, r, 4, f, "$M", bold=ends)
    put(BR, r, 5, None if ends else f"=IFERROR(D{r}/$D$17,0)", "%")
    put(BR, r, 6, f"=_xlfn.FORMULATEXT(D{r})")
checks = [
    ("Check", "Exit enterprise value (direct)", "Exit EBITDA x exit multiple", f"={V}B12*{A}B14", "$M"),
    ("Check", "Exit net debt (direct)", "Entry net debt + add-on debt - cumulative FCF", f"={A}B10+{V}B9-{V}D29", "$M"),
    ("Check", "Exit equity value (direct)", "Exit EV - exit net debt", "=D13-D14", "$M"),
    ("Check", "Bridge minus direct (must be 0)", "Reconciles the additive bridge to the direct calculation", "=ROUND(D12-D15,6)", "$M"),
    ("Result", "Total value created", "Exit equity - entry equity", "=D12-D4", "$M"),
    ("Result", "Gross MOIC", "Exit equity / entry equity", "=D12/D4", "x"),
    ("Result", "Gross IRR", "MOIC^(1/hold) - 1 (single entry and exit cash flow)", f"=D18^(1/{A}B13)-1", "%"),
]
for i, (s, lev, logic, f, fm) in enumerate(checks):
    r = 13 + i
    put(BR, r, 1, s, bold=True); put(BR, r, 2, lev); put(BR, r, 3, logic); put(BR, r, 4, f, fm)
    put(BR, r, 6, f"=_xlfn.FORMULATEXT(D{r})")
put(BR, 16, 5, '=IF(ABS(D16)<0.001,"OK - bridge reconciles","ERROR - bridge does not reconcile")', kind="k")
widths(BR, [9, 36, 56, 14, 22, 72]); BR.freeze_panes = "A4"

# ---------------- Lever Impact Ranking ----------------
title(LR, "Lever Impact Ranking", "Ranks bridge levers by equity value contribution and tests reliance on market vs operational levers.", 6)
header(LR, 3, ["Lever", "Category", "Value ($M)", "% of value created", "Rank (1 = largest)", "Controllability"])
levers = [("Operational", "High - pricing, sales capacity, retention"),
          ("Operational", "Medium - depends on signing and closing add-ons"),
          ("Operational", "High - procurement, labor productivity, pricing"),
          ("Market", "Low - exit market conditions"),
          ("Financial", "Medium - negotiated add-on purchase multiple"),
          ("Financial", "High - follows cash conversion and sweep"),
          ("Financial", "High - follows cash conversion")]
for i, (cat, ctl) in enumerate(levers):
    r = 4 + i; br = 5 + i
    put(LR, r, 1, f"={B}B{br}", bold=True); put(LR, r, 2, cat)
    put(LR, r, 3, f"={B}D{br}", "$M"); put(LR, r, 4, f"={B}E{br}", "%")
    put(LR, r, 5, f"=RANK(C{r},$C$4:$C$10,0)", "n"); put(LR, r, 6, ctl)
section(LR, 12, "Lever reliance tests")
header(LR, 13, ["Test", "Value", "Threshold", "Status"])
tests = [
    ("Operational levers ($M)", '=SUMIF(B4:B10,"Operational",C4:C10)', "$M", None, None),
    ("Market levers ($M)", '=SUMIF(B4:B10,"Market",C4:C10)', "$M", None, None),
    ("Financial levers ($M)", '=SUMIF(B4:B10,"Financial",C4:C10)', "$M", None, None),
    ("Operational share of value created", f"=IFERROR(B14/{B}D17,0)", "%", f"={A}B21", '=IF(B17>=C17,"PASS - at or above minimum","FAIL - below minimum")'),
    ("Multiple-change share of value created", f"=IFERROR(B15/{B}D17,0)", "%", f"={A}B20", '=IF(B18<=C18,"PASS - at or below maximum","FAIL - above maximum")'),
    ("Largest lever", "=INDEX(A4:A10,MATCH(1,E4:E10,0))", None, None, None),
    ("Largest lever share of value created", "=INDEX(D4:D10,MATCH(1,E4:E10,0))", "%", None, None),
]
for i, (lab, f, fm, th, st) in enumerate(tests):
    r = 14 + i
    put(LR, r, 1, lab, bold=True); put(LR, r, 2, f, fm)
    put(LR, r, 3, th, fm); put(LR, r, 4, st)
widths(LR, [40, 16, 14, 30, 16, 46]); LR.freeze_panes = "A4"

# ---------------- Summary Recommendation ----------------
title(SR, "Summary Recommendation", "Gates base-case returns, lever reliance and a downside case into one underwriting recommendation.", 3)
header(SR, 3, ["Metric", "Value", "Note"])
srows = [
    ("Entry equity value", f"={B}D4", "$M", "Bridge D4"),
    ("Exit equity value", f"={B}D12", "$M", "Bridge D12"),
    ("Total value created", f"={B}D17", "$M", "Bridge D17"),
    ("Gross MOIC", f"={B}D18", "x", "Bridge D18"),
    ("Gross IRR", f"={B}D19", "%", "Bridge D19"),
    ("Target gross MOIC", f"={A}B16", "x", "Input"),
    ("Target gross IRR", f"={A}B17", "%", "Input"),
    ("Operational share of value created", f"={L}B17", "%", "Lever Impact Ranking B17"),
    ("Minimum operational share", f"={A}B21", "%", "Input"),
    ("Multiple-change share of value created", f"={L}B18", "%", "Lever Impact Ranking B18"),
    ("Maximum multiple-change share", f"={A}B20", "%", "Input"),
    None,
    ("Downside organic exit revenue", f"={A}B5*(1+{V}B4*{A}B23)^{A}B13", "$M", "Organic CAGR x growth credited"),
    ("Downside exit EBITDA margin", f"={A}B6+({V}B11-{A}B6)*{A}B24", "%", "Entry margin + credited margin change"),
    ("Downside exit EBITDA", f"=(B16+{V}B6)*B17", "$M", "Signed add-on revenue retained"),
    ("Downside exit multiple", f"={A}B14-{A}B22", "x", "Underwritten exit multiple less haircut"),
    ("Downside cumulative FCF", f"={V}B14*({A}B13*{A}B7+(B18-{A}B7)*({A}B13+1)/2)", "$M", "Same linear EBITDA path as the schedule"),
    ("Downside exit net debt", f"={A}B10+{V}B9-B20", "$M", "Entry net debt + add-on debt - FCF"),
    ("Downside exit equity value", "=B18*B19-B21", "$M", "Downside EV - downside net debt"),
    ("Downside MOIC", "=B22/B4", "x", "Downside exit equity / entry equity"),
    ("Downside MOIC floor", f"={A}B19", "x", "Input"),
    ("Base-case MOIC floor", f"={A}B18", "x", "Input"),
    ("Bridge reconciliation", f"={B}E16", None, "Bridge E16"),
]
r = 4
for item in srows:
    if item is None:
        section(SR, r, "Downside case (stress test)"); r += 1; continue
    lab, f, fm, note = item
    put(SR, r, 1, lab, bold=True); put(SR, r, 2, f, fm); put(SR, r, 3, note); r += 1
assert r == 27
put(SR, 27, 1, "RECOMMENDATION", bold=True)
put(SR, 27, 2, '=IF(B26<>"OK - bridge reconciles","CHECK MODEL - BRIDGE DOES NOT RECONCILE",'
               'IF(OR(B7<B25,B23<B24),"DO NOT UNDERWRITE AT CURRENT TERMS",'
               'IF(AND(B7>=B9,B8>=B10,B11>=B12,B13<=B14),"UNDERWRITE - VALUE CREATION PLAN SUPPORTS RETURNS",'
               '"UNDERWRITE WITH CONDITIONS - REBALANCE LEVER RELIANCE")))', kind="k", bold=True)
put(SR, 27, 3, "Order: reconciliation -> floors (base and downside) -> targets and lever-reliance tests.")
put(SR, 28, 1, "Reasoning", bold=True)
put(SR, 28, 2, '="Base MOIC "&TEXT(B7,"0.00")&"x vs target "&TEXT(B9,"0.00")&"x; IRR "&TEXT(B8,"0.0%")&" vs "&TEXT(B10,"0.0%")'
               '&"; operational share "&TEXT(B11,"0%")&" (min "&TEXT(B12,"0%")&"); multiple share "&TEXT(B13,"0%")'
               '&" (max "&TEXT(B14,"0%")&"); downside MOIC "&TEXT(B23,"0.00")&"x (floor "&TEXT(B24,"0.00")&"x)."')
put(SR, 29, 1, "Recommendation formula (FORMULATEXT of B27)", bold=True)
put(SR, 29, 2, "=_xlfn.FORMULATEXT(B27)")
widths(SR, [44, 70, 46]); SR.freeze_panes = "A4"

# ---------------- Anonymized Worked Example ----------------
title(WE, "Anonymized Worked Example - HoldCo B", "Regional industrial services platform (code name). Figures scaled and illustrative; every output is a live link to the model tabs.", 6)
section(WE, 3, "Deal snapshot")
header(WE, 4, ["Item", "Value", "Unit"])
snap = [("Entry LTM EBITDA", f"={A}B7", "$M"), ("Entry EV / EBITDA multiple", f"={A}B8", "x"),
        ("Entry equity value", f"={A}B12", "$M"), ("Exit EBITDA", f"={V}B12", "$M"),
        ("Underwritten exit multiple", f"={A}B14", "x"), ("Hold period", f"={A}B13", "yrs")]
for i, (lab, f, u) in enumerate(snap):
    put(WE, 5 + i, 1, lab, bold=True); put(WE, 5 + i, 2, f, u); put(WE, 5 + i, 3, u)
section(WE, 12, "Bridge output: entry equity -> exit equity ($M)")
header(WE, 13, ["Bridge step", "Value ($M)", "% of value created"])
for i in range(9):
    r = 14 + i
    put(WE, r, 1, f"={B}B{4 + i}", bold=i in (0, 8)); put(WE, r, 2, f"={B}D{4 + i}", "$M", bold=i in (0, 8))
    put(WE, r, 3, None if i in (0, 8) else f"={B}E{4 + i}", "%")
section(WE, 24, "Named judgment calls")
header(WE, 25, ["ID", "Judgment call", "Underwritten choice", "Alternative rejected", "Exit equity excluded vs alternative ($M)", "Rationale"])
jcs = [
    ("JC1", "Exit multiple held at entry multiple", f'=TEXT({A}B14,"0.0")&"x exit multiple"', f'=TEXT({A}B15,"0.0")&"x management comp set"',
     f"={V}B12*({A}B15-{A}B14)", "No multiple expansion underwritten: LMM services exits are cyclical and the 9.5x comp set includes larger, more diversified peers."),
    ("JC2", "Margin expansion haircut", f'="+"&TEXT(({V}B11-{A}B6)*10000,"0")&" bps credited"', f'="+"&TEXT(({V}B16-{A}B6)*10000,"0")&" bps management plan"',
     f"={V}B10*({V}B16-{V}B11)*{A}B14", "Only procurement and pricing actions already signed are credited; the labor-productivity half of the plan is unproven."),
    ("JC3", "Unsigned add-on pipeline excluded", f'="$"&TEXT({V}B6,"0.0")&"M signed-LOI add-on"', f'="+$"&TEXT({V}B18,"0.0")&"M unsigned pipeline"',
     f"={V}B18*{V}B7*({A}B14-{V}B8)", "Unsigned targets are not underwritten; excluded value is shown net of purchase price at the same add-on multiple."),
    ("JC4", "Cash conversion haircut", f'=TEXT({V}B14,"0%")&" FCF / EBITDA"', f'=TEXT({V}B17,"0%")&" management case"',
     f"=({V}B17-{V}B14)*{V}C29", "Fleet capex catch-up in years 1-2 and higher cash taxes after the recap reduce free cash flow."),
]
for i, row in enumerate(jcs):
    r = 26 + i
    for c, v in enumerate(row, 1):
        put(WE, r, c, v, "$M" if c == 5 else None, bold=c == 1)
put(WE, 30, 1, "Total", bold=True); put(WE, 30, 2, "Equity value excluded by judgment calls (approximate; effects interact, not strictly additive)")
put(WE, 30, 5, "=SUM(E26:E29)", "$M", bold=True)
section(WE, 32, "Result")
res = [("Gross MOIC", f"={S}B7", "x"), ("Gross IRR", f"={S}B8", "%"), ("Downside MOIC", f"={S}B23", "x"),
       ("Largest lever", f"={L}B19", None), ("RECOMMENDATION", f"={S}B27", None), ("Reasoning", f"={S}B28", None)]
for i, (lab, f, fm) in enumerate(res):
    put(WE, 33 + i, 1, lab, bold=True); put(WE, 33 + i, 2, f, fm, kind="k" if lab == "RECOMMENDATION" else None, bold=lab == "RECOMMENDATION")
section(WE, 40, "How to read this example")
notes = ["1. Value created comes mostly from operational levers (organic growth, the signed add-on and +200 bps margin) and from deleveraging. No multiple expansion is underwritten.",
         "2. The judgment calls are conservative: reversing them would add the value in E26:E29. That upside is real but unproven, so it is kept out of the underwriting case.",
         "3. The downside case (Summary Recommendation rows 16-23) cuts the exit multiple by 1.0x, halves organic growth and removes the margin gain. It is still tested against the downside MOIC floor."]
for i, t in enumerate(notes):
    WE.cell(41 + i, 1, t).alignment = Alignment(wrap_text=False)
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Equity value bridge ($M)"; ch.legend = None
ch.add_data(Reference(WE, min_col=2, min_row=14, max_row=22)); ch.set_categories(Reference(WE, min_col=1, min_row=14, max_row=22))
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.tickLblPos = "low"; ch.height = 9; ch.width = 18
WE.add_chart(ch, "H3")
widths(WE, [36, 64, 28, 30, 22, 70])

# ---------------- ReadMe - Validator Instructions ----------------
title(RM, "Equity Value Creation Bridge Toolkit - v1.0", "Post Fiat Analytical Infrastructure Library | Reusable LMM private equity workflow | Formula-driven, anonymized", 5)
section(RM, 4, "Purpose")
RM["A5"] = ("Bridges sponsor entry equity to exit equity across seven levers: organic revenue growth, add-on revenue, margin change, "
            "multiple change, add-on purchase price, debt paydown and cash generation. It then ranks the levers, tests reliance on multiple expansion, "
            "runs a downside case and returns one underwriting recommendation.")
section(RM, 7, "How to use")
for i, t in enumerate(["1. Enter deal inputs (yellow cells) on Entry and Exit Assumptions B5:B24 and Value Driver Inputs B4:B18.",
                       "2. Check that Equity Value Creation Bridge E16 reads 'OK - bridge reconciles' and Value Driver Inputs B30 reads 'OK'.",
                       "3. Read the lever ranking and reliance tests, then the recommendation in Summary Recommendation B27 and its reasoning in B28.",
                       "4. Record each non-obvious underwriting choice as a named judgment call (see Anonymized Worked Example rows 26-29)."]):
    RM.cell(8 + i, 1, t)
section(RM, 13, "Tabs")
header(RM, 14, ["#", "Tab", "What it does"])
desc = ["Purpose, usage and validator checklist", "Entry valuation, exit multiple, return targets and downside settings",
        "Operating levers and the annual EBITDA / FCF schedule", "Seven-step additive bridge with reconciliation and FORMULATEXT",
        "Lever ranking, categories and reliance tests", "Base and downside returns gated into one recommendation",
        "HoldCo B example: bridge output, chart, 4 named judgment calls, result"]
for i, (n, t) in enumerate(zip(N, desc)):
    put(RM, 15 + i, 1, i + 1); put(RM, 15 + i, 2, n, bold=True); put(RM, 15 + i, 3, t)
RM["A23"] = "Colour legend: yellow = input, blue = formula, green = key output. Tab 1 is named 'ReadMe - Validator Instructions' because Excel does not allow '/' in tab names."
section(RM, 25, "Validator checklist (criterion -> tab -> cell range -> live check)")
header(RM, 26, ["#", "Criterion", "Tab", "Cell range", "Live check"])
chk = [
    ("Bridge from entry equity to exit equity, all live formulas", N[3], "D4:D12; FORMULATEXT in F4:F19", f"={B}E16"),
    ("Levers: EBITDA growth (organic + add-on), margin, multiple, debt paydown, cash generation", N[3], "B5:E11", f'=COUNTA({B}B5:B11)&" levers"'),
    ("Bridge reconciles to directly computed exit equity", N[3], "D13:D16; status E16", f'=TEXT({B}D16,"0.000000")'),
    ("Annual EBITDA / FCF schedule covers the hold period", N[2], "A22:F30; status B30", f"={V}B30"),
    ("Lever ranking and reliance tests", N[4], "A4:F10; A14:D20", f'={L}D17&" | "&{L}D18'),
    ("Clear recommendation with visible logic", N[5], "B27 (formula text in B29); reasoning B28", f"={S}B27"),
    ("Anonymized worked example with visible bridge output", N[6], "A14:C22 plus bar chart", f'="Exit equity $"&TEXT({W}B22,"0.00")&"M"'),
    ("At least 3 named judgment calls", N[6], "A26:F29", f'=COUNTA({W}A26:A29)&" judgment calls"'),
    ("Zero formula errors", "All tabs", "Excel SpecialCells(xlFormulas, xlErrors) scan", "See README.md (recalculated in Excel)"),
]
for i, row in enumerate(chk):
    r = 27 + i
    put(RM, r, 1, i + 1)
    for c, v in enumerate(row, 2):
        put(RM, r, c, v, bold=c == 2)
widths(RM, [6, 60, 34, 44, 40])

wb.save(OUT)
print("saved", OUT)


